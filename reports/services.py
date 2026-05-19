"""
Report generation services for the CRM system.
Handles PDF generation using ReportLab.
"""
import io
from datetime import datetime

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
    HRFlowable,
)
from reportlab.lib.enums import TA_CENTER, TA_LEFT


# ─── Color palette ────────────────────────────────────────────────────────────
DARK_BLUE = colors.HexColor('#1e3a5f')
TEAL_BLUE = colors.HexColor('#2c6e8a')
LIGHT_GRAY = colors.HexColor('#f3f4f6')
MID_GRAY = colors.HexColor('#9ca3af')
WHITE = colors.white


# ─── Report type metadata ─────────────────────────────────────────────────────

REPORT_CONFIG = {
    'clients': {
        'title': 'Reporte de Clientes',
        'columns': ['Nombre / Empresa', 'Correo Electrónico', 'Teléfono', 'Estado', 'Fecha de Registro'],
        'col_widths': [5.5 * cm, 5.5 * cm, 3.5 * cm, 2.5 * cm, 3.5 * cm],
        'landscape': False,
    },
    'sales': {
        'title': 'Reporte de Ventas (Oportunidades)',
        'columns': ['Cliente', 'Monto Estimado', 'Estado', 'Vendedor', 'Fecha Cierre Esperada'],
        'col_widths': [5 * cm, 3.5 * cm, 3.5 * cm, 4 * cm, 4 * cm],
        'landscape': False,
    },
    'opportunities': {
        'title': 'Reporte de Oportunidades',
        'columns': ['Título', 'Cliente', 'Valor Estimado', 'Etapa', 'Vendedor', 'Fecha Cierre'],
        'col_widths': [4 * cm, 3.5 * cm, 3 * cm, 3 * cm, 3.5 * cm, 3 * cm],
        'landscape': True,
    },
    'followups': {
        'title': 'Reporte de Seguimientos',
        'columns': ['Cliente', 'Tipo', 'Notas', 'Vendedor', 'Fecha'],
        'col_widths': [4 * cm, 2.5 * cm, 6 * cm, 4 * cm, 3.5 * cm],
        'landscape': False,
    },
}


# ─── Style helpers ────────────────────────────────────────────────────────────

def _get_styles():
    """Return a dict of custom paragraph styles."""
    base = getSampleStyleSheet()

    return {
        'title': ParagraphStyle(
            'CRMTitle',
            parent=base['Title'],
            fontSize=20,
            textColor=DARK_BLUE,
            spaceAfter=4,
            alignment=TA_CENTER,
            fontName='Helvetica-Bold',
        ),
        'subtitle': ParagraphStyle(
            'CRMSubtitle',
            parent=base['Normal'],
            fontSize=11,
            textColor=TEAL_BLUE,
            spaceAfter=2,
            alignment=TA_CENTER,
            fontName='Helvetica',
        ),
        'meta': ParagraphStyle(
            'CRMMeta',
            parent=base['Normal'],
            fontSize=9,
            textColor=colors.HexColor('#4a4a4a'),
            spaceAfter=2,
            alignment=TA_CENTER,
            fontName='Helvetica',
        ),
        'section': ParagraphStyle(
            'CRMSection',
            parent=base['Normal'],
            fontSize=11,
            textColor=DARK_BLUE,
            spaceBefore=10,
            spaceAfter=4,
            fontName='Helvetica-Bold',
        ),
        'stat_label': ParagraphStyle(
            'CRMStatLabel',
            parent=base['Normal'],
            fontSize=9,
            textColor=colors.HexColor('#4a4a4a'),
            fontName='Helvetica',
        ),
    }


def _build_table_style(num_rows):
    """
    Build a TableStyle with:
    - Dark header row (DARK_BLUE bg, white text)
    - Alternating row colors (white / light gray)
    - Borders and padding
    """
    commands = [
        # Header
        ('BACKGROUND', (0, 0), (-1, 0), DARK_BLUE),
        ('TEXTCOLOR', (0, 0), (-1, 0), WHITE),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        # All cells
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        # Borders
        ('BOX', (0, 0), (-1, -1), 0.5, MID_GRAY),
        ('INNERGRID', (0, 0), (-1, -1), 0.25, MID_GRAY),
        # Data rows
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ALIGN', (0, 1), (-1, -1), 'LEFT'),
    ]

    # Alternating row colors
    for row in range(1, num_rows):
        bg = WHITE if row % 2 == 1 else LIGHT_GRAY
        commands.append(('BACKGROUND', (0, row), (-1, row), bg))

    return TableStyle(commands)


def _add_page_number(canvas, doc):
    """Draw footer with page number and generation timestamp on every page."""
    canvas.saveState()
    width, _ = doc.pagesize
    timestamp = datetime.now().strftime('%d/%m/%Y %H:%M:%S')
    footer_text = f"Generado el {timestamp}  |  Página {doc.page}"
    canvas.setFont('Helvetica', 7)
    canvas.setFillColor(MID_GRAY)
    canvas.drawCentredString(width / 2, 1.0 * cm, footer_text)
    canvas.setStrokeColor(MID_GRAY)
    canvas.setLineWidth(0.3)
    canvas.line(doc.leftMargin, 1.4 * cm, width - doc.rightMargin, 1.4 * cm)
    canvas.restoreState()


# ─── Row builders ─────────────────────────────────────────────────────────────

_STAGE_LABELS = {
    'prospeccion': 'Prospección',
    'calificacion': 'Calificación',
    'propuesta': 'Propuesta',
    'negociacion': 'Negociación',
    'cierre_ganado': 'Cierre Ganado',
    'cierre_perdido': 'Cierre Perdido',
}

_FOLLOW_UP_LABELS = {
    'call': 'Llamada',
    'email': 'Email',
    'meeting': 'Reunión',
    'other': 'Otro',
}


def _rows_and_summary_for_clients(queryset):
    total = queryset.count()
    active = queryset.filter(is_active=True).count()
    summary = {
        'Total de clientes': total,
        'Clientes activos': active,
        'Clientes inactivos': total - active,
    }
    rows = []
    for c in queryset:
        rows.append([
            c.company_name,
            c.email,
            c.phone,
            'Activo' if c.is_active else 'Inactivo',
            c.created_at.strftime('%d/%m/%Y') if c.created_at else '—',
        ])
    return rows, summary


def _rows_and_summary_for_sales(queryset):
    total = queryset.count()
    total_value = sum(o.estimated_value for o in queryset)
    won = queryset.filter(stage='cierre_ganado').count()
    lost = queryset.filter(stage='cierre_perdido').count()
    summary = {
        'Total de oportunidades': total,
        'Valor total estimado': f"${total_value:,.2f}",
        'Cerradas ganadas': won,
        'Cerradas perdidas': lost,
    }
    rows = []
    for o in queryset:
        vendedor = (
            o.assigned_vendedor.get_full_name() or o.assigned_vendedor.username
            if o.assigned_vendedor else '—'
        )
        rows.append([
            o.client.company_name if o.client else '—',
            f"${o.estimated_value:,.2f}",
            _STAGE_LABELS.get(o.stage, o.stage),
            vendedor,
            o.expected_close_date.strftime('%d/%m/%Y') if o.expected_close_date else '—',
        ])
    return rows, summary


def _rows_and_summary_for_opportunities(queryset):
    total = queryset.count()
    total_value = sum(o.estimated_value for o in queryset)
    total_weighted = sum(o.weighted_value for o in queryset)
    summary = {
        'Total de oportunidades': total,
        'Valor total estimado': f"${total_value:,.2f}",
        'Valor ponderado total': f"${total_weighted:,.2f}",
    }
    rows = []
    for o in queryset:
        vendedor = (
            o.assigned_vendedor.get_full_name() or o.assigned_vendedor.username
            if o.assigned_vendedor else '—'
        )
        rows.append([
            o.title,
            o.client.company_name if o.client else '—',
            f"${o.estimated_value:,.2f}",
            _STAGE_LABELS.get(o.stage, o.stage),
            vendedor,
            o.expected_close_date.strftime('%d/%m/%Y') if o.expected_close_date else '—',
        ])
    return rows, summary


def _rows_and_summary_for_followups(queryset):
    total = queryset.count()
    summary = {
        'Total de seguimientos': total,
        'Llamadas': queryset.filter(follow_up_type='call').count(),
        'Emails': queryset.filter(follow_up_type='email').count(),
        'Reuniones': queryset.filter(follow_up_type='meeting').count(),
    }
    rows = []
    for f in queryset:
        client_name = '—'
        if f.client:
            client_name = f.client.company_name
        elif f.opportunity and f.opportunity.client:
            client_name = f.opportunity.client.company_name

        creator = (
            f.created_by.get_full_name() or f.created_by.username
            if f.created_by else '—'
        )
        notes = (f.notes[:60] + '…') if len(f.notes) > 60 else f.notes
        rows.append([
            client_name,
            _FOLLOW_UP_LABELS.get(f.follow_up_type, f.follow_up_type),
            notes,
            creator,
            f.date.strftime('%d/%m/%Y %H:%M') if f.date else '—',
        ])
    return rows, summary


_ROW_BUILDERS = {
    'clients': _rows_and_summary_for_clients,
    'sales': _rows_and_summary_for_sales,
    'opportunities': _rows_and_summary_for_opportunities,
    'followups': _rows_and_summary_for_followups,
}


# ─── Main PDF generator ────────────────────────────────────────────────────────

def generate_pdf_report(queryset, report_type: str, filters: dict, user) -> bytes:
    """
    Generate a PDF report and return the raw bytes.

    Args:
        queryset:    Pre-filtered queryset of records to include.
        report_type: One of 'clients', 'sales', 'opportunities', 'followups'.
        filters:     Dict with keys date_from, date_to, vendedor, status, industry.
        user:        The authenticated Django user requesting the report.

    Returns:
        bytes: The PDF file content.

    Raises:
        ValueError: If report_type is not recognised.
    """
    if report_type not in REPORT_CONFIG:
        raise ValueError(f"Tipo de reporte no válido: {report_type}")

    config = REPORT_CONFIG[report_type]
    styles = _get_styles()

    # Build rows and summary stats
    builder = _ROW_BUILDERS[report_type]
    data_rows, summary = builder(queryset)

    # ── Document setup ──────────────────────────────────────────────────────
    buffer = io.BytesIO()
    page_size = landscape(A4) if config['landscape'] else A4
    doc = SimpleDocTemplate(
        buffer,
        pagesize=page_size,
        leftMargin=1.8 * cm,
        rightMargin=1.8 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title=config['title'],
        author=user.get_full_name() or user.username,
    )

    story = []

    # ── Header ──────────────────────────────────────────────────────────────
    story.append(Paragraph('CRM — Sistema de Gestión de Ventas', styles['title']))
    story.append(Paragraph(config['title'], styles['subtitle']))
    story.append(Spacer(1, 0.2 * cm))

    # Date range
    date_from = filters.get('date_from')
    date_to = filters.get('date_to')
    if date_from and date_to:
        date_range = f"Período: {date_from} — {date_to}"
    elif date_from:
        date_range = f"Desde: {date_from}"
    elif date_to:
        date_range = f"Hasta: {date_to}"
    else:
        date_range = "Período: Todos los registros"

    story.append(Paragraph(date_range, styles['meta']))
    story.append(Paragraph(
        f"Generado por: {user.get_full_name() or user.username}"
        f"  |  {datetime.now().strftime('%d/%m/%Y %H:%M')}",
        styles['meta'],
    ))
    story.append(Spacer(1, 0.3 * cm))
    story.append(HRFlowable(width='100%', thickness=1.5, color=DARK_BLUE, spaceAfter=8))

    # ── Summary stats ────────────────────────────────────────────────────────
    if summary:
        story.append(Paragraph('Resumen', styles['section']))
        stat_items = list(summary.items())
        stat_rows = []
        for i in range(0, len(stat_items), 2):
            left_label, left_val = stat_items[i]
            if i + 1 < len(stat_items):
                right_label, right_val = stat_items[i + 1]
            else:
                right_label, right_val = '', ''
            stat_rows.append([
                Paragraph(f"<b>{left_label}:</b> {left_val}", styles['stat_label']),
                Paragraph(
                    f"<b>{right_label}:</b> {right_val}" if right_label else '',
                    styles['stat_label'],
                ),
            ])

        stat_table = Table(stat_rows, colWidths=['50%', '50%'])
        stat_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), LIGHT_GRAY),
            ('BOX', (0, 0), (-1, -1), 0.5, MID_GRAY),
            ('INNERGRID', (0, 0), (-1, -1), 0.25, MID_GRAY),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('RIGHTPADDING', (0, 0), (-1, -1), 8),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        story.append(stat_table)
        story.append(Spacer(1, 0.4 * cm))

    # ── Data table ───────────────────────────────────────────────────────────
    story.append(Paragraph('Detalle de Registros', styles['section']))

    if data_rows:
        table_data = [config['columns']] + data_rows
        data_table = Table(table_data, colWidths=config['col_widths'], repeatRows=1)
        data_table.setStyle(_build_table_style(len(table_data)))
        story.append(data_table)
    else:
        story.append(Paragraph(
            'No se encontraron registros con los filtros aplicados.',
            styles['meta'],
        ))

    # ── Build PDF ────────────────────────────────────────────────────────────
    doc.build(story, onFirstPage=_add_page_number, onLaterPages=_add_page_number)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes


# ─── Excel export ──────────────────────────────────────────────────────────────

import csv as _csv

import openpyxl
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Border,
    Side,
)
from openpyxl.utils import get_column_letter

# Hex colours for Excel formatting (no leading #)
_XL_DARK_BLUE = "1E3A5F"
_XL_MEDIUM_BLUE = "2C6E8A"
_XL_ALT_ROW = "EBF4FA"
_XL_SUMMARY_ROW = "D9E8F5"

# Column definitions per report type
_EXCEL_COLUMNS = {
    "clients": {
        "title": "Reporte de Clientes",
        "headers": ["Empresa", "Email", "Telefono", "Estado", "Ciudad", "Fecha de Creacion"],
    },
    "sales": {
        "title": "Reporte de Ventas",
        "headers": ["Cliente", "Monto", "Estado", "Vendedor", "Fecha", "Notas"],
    },
    "opportunities": {
        "title": "Reporte de Oportunidades",
        "headers": ["Titulo", "Cliente", "Valor Estimado", "Etapa", "Vendedor", "Fecha Esperada de Cierre"],
    },
    "followups": {
        "title": "Reporte de Seguimientos",
        "headers": ["Cliente", "Tipo de Contacto", "Notas", "Vendedor", "Fecha"],
    },
}


def _queryset_to_export_rows(report_type, queryset):
    """Convert a queryset to a list of plain-string rows for Excel/CSV output."""
    rows = []
    stage_labels = {
        "prospeccion": "Prospeccion",
        "calificacion": "Calificacion",
        "propuesta": "Propuesta",
        "negociacion": "Negociacion",
        "cierre_ganado": "Cierre Ganado",
        "cierre_perdido": "Cierre Perdido",
    }
    type_labels = {"call": "Llamada", "email": "Email", "meeting": "Reunion", "other": "Otro"}

    if report_type == "clients":
        for c in queryset:
            rows.append([
                c.company_name,
                c.email,
                c.phone,
                "Activo" if c.is_active else "Inactivo",
                c.address or "",
                c.created_at.strftime("%Y-%m-%d") if c.created_at else "",
            ])

    elif report_type in ("sales", "opportunities"):
        for o in queryset:
            client_name = o.client.company_name if o.client else ""
            vendedor_name = o.assigned_vendedor.get_full_name() if o.assigned_vendedor else ""
            stage_display = stage_labels.get(o.stage, o.stage)
            if report_type == "sales":
                rows.append([
                    client_name,
                    str(o.estimated_value),
                    stage_display,
                    vendedor_name,
                    o.created_at.strftime("%Y-%m-%d") if o.created_at else "",
                    o.loss_reason or "",
                ])
            else:
                rows.append([
                    o.title,
                    client_name,
                    str(o.estimated_value),
                    stage_display,
                    vendedor_name,
                    o.expected_close_date.strftime("%Y-%m-%d") if o.expected_close_date else "",
                ])

    elif report_type == "followups":
        for f in queryset:
            client_name = ""
            if f.client:
                client_name = f.client.company_name
            elif f.opportunity and f.opportunity.client:
                client_name = f.opportunity.client.company_name
            rows.append([
                client_name,
                type_labels.get(f.follow_up_type, f.follow_up_type),
                f.notes,
                f.created_by.get_full_name() if f.created_by else "",
                f.date.strftime("%Y-%m-%d %H:%M") if f.date else "",
            ])

    return rows


def generate_excel_export(queryset, report_type: str, filters: dict) -> bytes:
    """
    Generate an Excel (.xlsx) report from a queryset and return the raw bytes.

    Layout:
        Row 1  - Report title (merged, dark blue bg, white bold text)
        Row 2  - Date range / filter info (merged, medium blue bg)
        Row 3  - Empty spacer
        Row 4  - Column headers (medium blue bg, bold white text)
        Rows 5+ - Data rows with alternating colours
        Last   - Total record count summary row

    Column widths are auto-fitted and the header row is frozen.

    Args:
        queryset: Django queryset already filtered and scoped.
        report_type (str): One of clients, sales, opportunities, followups.
        filters (dict): Applied filters (date_from, date_to, vendedor, status).

    Returns:
        bytes: Raw .xlsx file content.
    """
    config = _EXCEL_COLUMNS.get(report_type, {"title": "Reporte", "headers": []})
    title = config["title"]
    headers = config["headers"]
    rows = _queryset_to_export_rows(report_type, queryset)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]

    dark_blue_fill = PatternFill(fill_type="solid", fgColor=_XL_DARK_BLUE)
    medium_blue_fill = PatternFill(fill_type="solid", fgColor=_XL_MEDIUM_BLUE)
    alt_row_fill = PatternFill(fill_type="solid", fgColor=_XL_ALT_ROW)
    summary_fill = PatternFill(fill_type="solid", fgColor=_XL_SUMMARY_ROW)

    white_bold_font = Font(bold=True, color="FFFFFF", size=13)
    white_font = Font(color="FFFFFF", size=10)
    header_font = Font(bold=True, color="FFFFFF", size=11)
    normal_font = Font(size=10)
    summary_font = Font(bold=True, size=10)

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    num_cols = max(len(headers), 1)

    # Row 1: Title
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = white_bold_font
    title_cell.fill = dark_blue_fill
    title_cell.alignment = center_align

    # Row 2: Filter info
    ws.row_dimensions[2].height = 20
    today_str = datetime.now().strftime("%Y-%m-%d")
    date_from = filters.get("date_from")
    date_to = filters.get("date_to")
    filter_text = "Generado: " + today_str
    if date_from or date_to:
        filter_text += "  |  Periodo: " + str(date_from or "---") + " a " + str(date_to or "---")
    seller = filters.get("vendedor") or filters.get("seller")
    if seller:
        seller_name = seller.get_full_name() if hasattr(seller, "get_full_name") else str(seller)
        filter_text += "  |  Vendedor: " + seller_name
    status = filters.get("status")
    if status:
        filter_text += "  |  Estado/Etapa: " + str(status)

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols)
    info_cell = ws.cell(row=2, column=1, value=filter_text)
    info_cell.font = white_font
    info_cell.fill = medium_blue_fill
    info_cell.alignment = center_align

    # Row 3: Spacer
    ws.row_dimensions[3].height = 8

    # Row 4: Column headers
    ws.row_dimensions[4].height = 22
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = medium_blue_fill
        cell.alignment = center_align
        cell.border = thin_border

    # Rows 5+: Data rows with alternating colours
    for row_idx, row_data in enumerate(rows, start=5):
        ws.row_dimensions[row_idx].height = 18
        use_alt = (row_idx % 2 == 0)
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = normal_font
            if use_alt:
                cell.fill = alt_row_fill
            cell.alignment = left_align
            cell.border = thin_border

    # Last row: Summary
    if rows:
        summary_row = len(rows) + 5
        ws.row_dimensions[summary_row].height = 20
        ws.merge_cells(
            start_row=summary_row, start_column=1,
            end_row=summary_row, end_column=num_cols,
        )
        total_cell = ws.cell(
            row=summary_row, column=1,
            value="Total de registros: " + str(len(rows)),
        )
        total_cell.font = summary_font
        total_cell.fill = summary_fill
        total_cell.alignment = left_align

    # Auto-fit column widths
    for col_idx in range(1, num_cols + 1):
        col_letter = get_column_letter(col_idx)
        max_length = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 4, 50)

    # Freeze panes below header row
    ws.freeze_panes = "A5"

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


# ─── CSV export ────────────────────────────────────────────────────────────────


def generate_csv_export(queryset, report_type: str) -> str:
    """
    Generate a CSV report from a queryset and return the content as a string.

    The caller is responsible for prepending the UTF-8 BOM (b'\xef\xbb\xbf')
    before writing to the HTTP response so that Excel opens the file correctly.

    First row contains column headers; subsequent rows contain data.

    Args:
        queryset: Django queryset already filtered and scoped.
        report_type (str): One of clients, sales, opportunities, followups.

    Returns:
        str: CSV content as a Unicode string.
    """
    config = _EXCEL_COLUMNS.get(report_type, {"title": "Reporte", "headers": []})
    headers = config["headers"]
    rows = _queryset_to_export_rows(report_type, queryset)

    output = io.StringIO()
    writer = _csv.writer(output)

    if headers:
        writer.writerow(headers)

    for row in rows:
        writer.writerow(row)

    return output.getvalue()
